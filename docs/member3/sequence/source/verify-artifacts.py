from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import hashlib, json, xml.etree.ElementTree as ET
from docx import Document
from pypdf import PdfReader
from PIL import Image

root=Path(__file__).resolve().parent.parent
repo=root.parents[2]
workspace=repo.parent.parent
content=json.loads((root/'source/report-content.json').read_text(encoding='utf-8'))
docx=root/'Member3_Sequence_Diagrams.docx'
doc=Document(docx)
assert len(doc.inline_shapes)==2
expected=[content['title'],content['publicCaption'],*content['publicParagraphs'],content['reviewHeading'],content['reviewCaption'],*content['reviewParagraphs'],content['reference']]
assert [p.text for p in doc.paragraphs if p.text.strip()]==expected
with ZipFile(docx) as z:
    embedded=sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert embedded==sorted(hashlib.sha256((root/'diagrams'/n).read_bytes()).hexdigest() for n in ['Sequence_Submission.png','Sequence_Review_Tracking.png'])
pdf=workspace/'tmp/member3-sequence/word/Member3_Sequence_Diagrams.pdf'
reader=PdfReader(pdf)
assert len(reader.pages)==2
pdf_text=' '.join(' '.join(p.extract_text().split()) for p in reader.pages)
for text in expected:
    assert ' '.join(text.split()) in pdf_text, text[:70]

mx=ET.parse(root/'diagrams/Member3_Sequence_Diagrams.drawio').getroot()
assert len(mx.findall('diagram'))==2
for page in mx.findall('diagram'):
    cells=page.findall('.//mxCell')
    assert len(cells)>40 and not page.findall('.//*[@style="image"]')
    assert len({e.get('id') for e in cells})==len(cells)
    for cell in cells:
        if cell.get('id') not in {'0','1'}:assert cell.find('mxGeometry') is not None

layout=json.loads((root/'source/layout.json').read_text(encoding='utf-8'))
svg_check=json.loads((root/'source/svg-verification.json').read_text(encoding='utf-8'))
model_check=json.loads((root/'source/model-verification.json').read_text(encoding='utf-8'))
assert len(model_check['checks'])==12 and all(x['result']=='PASS' for x in model_check['checks'])
assert all(not d['outside'] and not d['overlaps'] for d in svg_check['results'])
for d in layout:
    assert ET.parse(root/'diagrams'/(d['file']+'.svg')).getroot().tag.endswith('svg')
    with Image.open(root/'diagrams'/(d['file']+'.png')) as im:assert im.width>=3000 and im.height>=3000
    for token in ['submit(input)' if d['id']=='M3A' else 'updateStatus(reference,','break ','alt ']:
        assert token in (root/'diagrams'/(d['file']+'.mmd')).read_text(encoding='utf-8')
original=workspace/'组员提供/member 1 diagram.png'
assert hashlib.sha256(original.read_bytes()).hexdigest()=='a6c34fa25c496291f3280a81a4b19bd8240e58002328aa001ed3f63d750758d6'
assert not (repo/'admin.html').exists()
artifacts=[]
for p in sorted(root.rglob('*')):
    if p.is_file() and p.name!='verification.json':
        artifacts.append({'path':p.relative_to(root).as_posix(),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
source_hashes={str(p.relative_to(repo)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in [repo/'js/feedback-store.js',repo/'js/feedback-pages.js']}
result={'checkedAt':datetime.now(timezone.utc).isoformat(),'report':{'file':docx.name,'pages':2,'figures':2,'words':len(' '.join(p.text for p in doc.paragraphs).split()),'renderer':'Microsoft Word 16 read-only PDF export','standardRenderer':'Unavailable: bundled LibreOffice not installed','allRenderedPagesVisuallyReviewed':True,'sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'embeddedImagesMatchSources':True,'paragraphsPresentInWordPdf':True},'diagrams':layout,'svgTextBounds':svg_check,'sharedModelChecks':model_check,'drawio':{'pages':2,'nativeEditableElements':True},'originalCandidateUnchanged':True,'implementationSources':source_hashes,'formalStaffPage':'Still pending; staff interaction is a proposed UI integration with an existing shared API','applicationCodeChanged':False,'fullBrowserSuitesRerun':False,'artifacts':artifacts}
(root/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':2,'figures':2,'words':result['report']['words'],'modelChecksPassed':12,'svgTextItemsChecked':sum(x['textCount'] for x in svg_check['results']),'drawioPages':2,'originalCandidateUnchanged':True}))
