from pathlib import Path
from datetime import datetime, timezone
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
data = json.loads((root/'source/report-content.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.6)
section.left_margin = section.right_margin = Inches(0.65)
for name in ['Normal','Title','Heading 1','Caption']:
    s=doc.styles[name]
    s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0);s.font.underline=False
    s.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial')
    s.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
    ppr=s.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in ppr.findall(qn('w:pBdr')):ppr.remove(border)
normal=doc.styles['Normal']
normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.03
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
title=doc.styles['Title'];title.font.size=Pt(18);title.font.bold=True
title.paragraph_format.space_after=Pt(6);title.paragraph_format.line_spacing=1
title.paragraph_format.keep_with_next=True
heading=doc.styles['Heading 1'];heading.font.size=Pt(13);heading.font.bold=True
heading.paragraph_format.space_before=Pt(0);heading.paragraph_format.space_after=Pt(6)
heading.paragraph_format.keep_with_next=True
caption=doc.styles['Caption'];caption.font.size=Pt(9);caption.font.italic=False
caption.paragraph_format.line_spacing=1;caption.paragraph_format.space_after=Pt(6)
caption.paragraph_format.keep_with_next=True

def figure(name,caption_text):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(root/'diagrams'/name),width=Inches(6.65))
    doc.inline_shapes[-1]._inline.docPr.set('descr',caption_text)
    doc.add_paragraph(caption_text,style='Caption')

doc.add_paragraph(data['title'],style='Title')
figure('Sequence_Submission.png',data['publicCaption'])
for text in data['publicParagraphs']:doc.add_paragraph(text)
p=doc.add_paragraph(data['reviewHeading'],style='Heading 1')
p.paragraph_format.page_break_before=True
figure('Sequence_Review_Tracking.png',data['reviewCaption'])
for text in data['reviewParagraphs']:doc.add_paragraph(text)
p=doc.add_paragraph(data['reference'])
for run in p.runs:run.font.size=Pt(8.5)
p.paragraph_format.space_after=Pt(0)
doc.core_properties.title=data['title']
doc.core_properties.subject='CMT322 Part A Member 3 sequence diagrams and explanation'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.comments=''
doc.core_properties.created=datetime.now(timezone.utc)
doc.core_properties.modified=datetime.now(timezone.utc)
target=root/'Member3_Sequence_Diagrams.docx'
doc.save(target)
lines=['# '+data['title'],'','![Figure M3A](diagrams/Sequence_Submission.png)','',data['publicCaption'],'']
lines += [p+'\n' for p in data['publicParagraphs']]
lines += ['## '+data['reviewHeading'],'','![Figure M3B](diagrams/Sequence_Review_Tracking.png)','',data['reviewCaption'],'']
lines += [p+'\n' for p in data['reviewParagraphs']]
lines += [data['reference'],'','Reference URL: https://www.omg.org/spec/UML/2.5.1','']
(root/'Member3_Sequence_Diagrams.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({'docx':target.name,'figures':len(doc.inline_shapes),'words':len(' '.join(p.text for p in doc.paragraphs).split())}))
