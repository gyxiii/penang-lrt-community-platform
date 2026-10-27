from pathlib import Path
import json
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
content = json.loads((root / 'source' / 'report-content.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.7)
section.header_distance = section.footer_distance = Inches(0.25)

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in ppr.findall(qn('w:pBdr')):
            ppr.remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.size = Pt(20)
title.font.bold = True
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(8)
title.paragraph_format.line_spacing = 1
title.paragraph_format.keep_with_next = True
heading = doc.styles['Heading 1']
heading.font.size = Pt(13)
heading.font.bold = True
heading.paragraph_format.space_before = Pt(8)
heading.paragraph_format.space_after = Pt(6)
heading.paragraph_format.keep_with_next = True
caption_style = doc.styles['Caption']
caption_style.font.size = Pt(9)
caption_style.font.bold = False
caption_style.font.italic = False
caption_style.paragraph_format.line_spacing = 1
caption_style.paragraph_format.space_before = Pt(4)
caption_style.paragraph_format.space_after = Pt(8)
caption_style.paragraph_format.keep_with_next = False

def paragraph(text):
    return doc.add_paragraph(text)

def figure(index, width=6.8677):
    item = content['figures'][index]
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(root / 'assets' / item['file']), width=Inches(width))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set('descr', item['caption'])
    cap = doc.add_paragraph(style='Caption')
    cap.add_run('Figure ' + item['id'] + '  ').bold = True
    cap.add_run(item['caption'])

doc.add_paragraph(content['title'], style='Title')
paragraph(content['intro'])
figure(0)
doc.add_paragraph(content['informationHeading'], style='Heading 1')
paragraph(content['information'])
paragraph(content['staff'])
figure(1)
p = doc.add_paragraph(content['feedbackHeading'], style='Heading 1')
p.paragraph_format.page_break_before = True
paragraph(content['feedback'])
figure(2, width=6.5)
paragraph(content['tracking'])
figure(3, width=6.5)
paragraph(content['scope'])

doc.core_properties.title = content['title']
doc.core_properties.subject = 'Prototype Design section for the CMT322 Part A group report'
doc.core_properties.author = 'MutiaraConnect Project Group'
doc.core_properties.last_modified_by = 'MutiaraConnect Project Group'
doc.core_properties.created = datetime.now(timezone.utc)
doc.core_properties.modified = datetime.now(timezone.utc)
doc.core_properties.keywords = 'CMT322, Part A, prototype design, feedback, user interface'
doc.core_properties.comments = ''
target = root / 'Member3_Prototype_Design.docx'
doc.save(target)

lines = ['# ' + content['title'], '', content['intro'], '']
def md_figure(index):
    item = content['figures'][index]
    return [f"![Figure {item['id']}](assets/{item['file']})", '', f"**Figure {item['id']}**  {item['caption']}", '']
lines += md_figure(0)
lines += ['## ' + content['informationHeading'], '', content['information'], '', content['staff'], '']
lines += md_figure(1)
lines += ['## ' + content['feedbackHeading'], '', content['feedback'], '']
lines += md_figure(2)
lines += [content['tracking'], '']
lines += md_figure(3)
lines += [content['scope'], '']
(root / 'Member3_Prototype_Design.md').write_text('\n'.join(lines), encoding='utf-8')
print(json.dumps({'docx': str(target), 'figures': len(doc.inline_shapes), 'paragraphs': len(doc.paragraphs), 'words': len(' '.join(p.text for p in doc.paragraphs).split())}))
